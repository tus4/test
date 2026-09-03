"""Workbook discovery and sheet extraction.

Workbooks are read with openpyxl in read-only mode with ``data_only=True`` so
formula cells yield their cached values (``None`` when the file was never
calculated by Excel). Sheets are streamed; only the first ``max_rows`` rows are
kept in memory, aggregates are computed over up to ``max_scan_rows`` rows.
"""

from __future__ import annotations

import fnmatch
import math
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from .analyze import (
    ColumnAcc,
    bucket_categories,
    convert_value,
    day_key,
    is_blank,
    is_error_text,
    is_placeholder,
    month_key,
    numeric_of,
    parse_numeric_text,
)
from .common import BuildError, fingerprint, truncate

WORKBOOK_SUFFIXES = (".xlsx", ".xlsm")
HARD_COLUMN_CAP = 1000          # columns tracked during the scan before the configured cap applies
MAX_HEADER_SEARCH_ROWS = 10000  # give up looking for a header after this many physical rows
MERGE_XML_MAX_BYTES = 50_000_000
MAX_LINE_POINTS = 60
MAX_DAY_BUCKETS = 400

_NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_NS_PKG = "{http://schemas.openxmlformats.org/package/2006/relationships}"

EXCLUSION_REASONS = {
    "config": "설정으로 제외한 열",
    "sensitive-name": "열 이름이 개인정보로 추정되어 자동 제외한 열 (공개하려면 allow_columns 에 추가)",
    "sensitive-values": "값이 이메일/전화번호/주민번호 형태여서 자동 제외한 열 (공개하려면 allow_columns 에 추가)",
}


# ---------------------------------------------------------------------------
# discovery
# ---------------------------------------------------------------------------
def path_matches(rel_path: str, patterns: Iterable[str]) -> bool:
    """Glob match on a POSIX relative path, case-insensitive (``Report.XLSX`` matches
    ``**/*.xlsx``, like discovery itself). ``**/x`` also matches ``x`` at the root."""
    lowered = rel_path.lower()
    for pat in patterns:
        pat = pat.lower()
        if fnmatch.fnmatchcase(lowered, pat):
            return True
        if pat.startswith("**/") and fnmatch.fnmatchcase(lowered, pat[3:]):
            return True
    return False


def name_matches(name: str, patterns: Iterable[str]) -> bool:
    lowered = name.lower()
    return any(fnmatch.fnmatchcase(lowered, pat.lower()) for pat in patterns)


def sensitive_name(name: str, patterns: Iterable[str]) -> bool:
    lowered = name.lower()
    return any(p.lower() in lowered for p in patterns)


def discover_workbooks_with_excluded(source: Path, cfg: Dict[str, Any]) -> Tuple[List[Tuple[Path, str]], int]:
    """(path-sorted list of (absolute path, relative POSIX path), number of workbooks excluded by config)."""
    found: List[Tuple[Path, str]] = []
    excluded = 0
    for path in source.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in WORKBOOK_SUFFIXES:
            continue
        rel = path.relative_to(source).as_posix()
        parts = rel.split("/")
        if any(part.startswith(".") for part in parts):
            continue  # .git, hidden folders
        if path.name.startswith("~$"):
            continue  # Excel lock files
        if cfg["include_files"] and not path_matches(rel, cfg["include_files"]):
            excluded += 1
            continue
        if path_matches(rel, cfg["exclude_files"]):
            excluded += 1
            continue
        found.append((path, rel))
    return sorted(found, key=lambda item: item[1]), excluded


def discover_workbooks(source: Path, cfg: Dict[str, Any]) -> List[Tuple[Path, str]]:
    """Return (absolute path, relative POSIX path) for every workbook to publish, path-sorted."""
    return discover_workbooks_with_excluded(source, cfg)[0]


# ---------------------------------------------------------------------------
# merged cells (best effort, header rows only)
# ---------------------------------------------------------------------------
def header_merged_ranges(path: Path, max_row: int) -> Dict[str, List[Tuple[int, int, int, int]]]:
    """Merged ranges (min_col, min_row, max_col, max_row) that start within the first
    ``max_row`` rows, keyed by sheet name. Parsed straight from the sheet XML because
    openpyxl's read-only mode does not expose merged cells. Returns {} on any problem."""
    try:
        with zipfile.ZipFile(path) as zf:
            wb_xml = ET.fromstring(zf.read("xl/workbook.xml"))
            rels_xml = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
            targets = {r.get("Id"): r.get("Target") for r in rels_xml.iter(_NS_PKG + "Relationship")}
            sheets = wb_xml.find(_NS_MAIN + "sheets")
            if sheets is None:
                return {}
            out: Dict[str, List[Tuple[int, int, int, int]]] = {}
            for sheet in sheets.iter(_NS_MAIN + "sheet"):
                name = sheet.get("name")
                target = targets.get(sheet.get(_NS_REL + "id"))
                if not name or not target:
                    continue
                xml_path = target.lstrip("/") if target.startswith("/") else "xl/" + target
                try:
                    info = zf.getinfo(xml_path)
                except KeyError:
                    continue
                if info.file_size > MERGE_XML_MAX_BYTES:
                    continue
                ranges: List[Tuple[int, int, int, int]] = []
                with zf.open(xml_path) as fh:
                    for _, element in ET.iterparse(fh, events=("end",)):
                        if element.tag == _NS_MAIN + "mergeCell":
                            ref = element.get("ref")
                            if ref:
                                try:
                                    c1, r1, c2, r2 = range_boundaries(ref)
                                except Exception:  # malformed ref: ignore
                                    continue
                                if r1 is not None and r1 <= max_row:
                                    ranges.append((int(c1), int(r1), int(c2), int(r2)))
                        element.clear()
                if ranges:
                    out[name] = sorted(ranges)
            return out
    except Exception:
        return {}


# ---------------------------------------------------------------------------
# header detection
# ---------------------------------------------------------------------------
def detect_header(window: List[List[Any]]) -> Tuple[int, str]:
    """Pick the header row inside ``window`` (non-blank rows only).

    Rule: the first row whose non-empty count reaches max(2, 50% of the widest
    row); among the first three such rows prefer one that is mostly text. Rows
    above it are joined into a title. Returns (index in window, title)."""
    counts = [sum(1 for v in row if not is_blank(v)) for row in window]
    width = max(counts)
    threshold = max(2, math.ceil(width * 0.5))
    candidates = [i for i, c in enumerate(counts) if c >= threshold]
    if candidates:
        header = candidates[0]
        for i in candidates[:3]:
            non_blank = [v for v in window[i] if not is_blank(v)]
            n_text = sum(1 for v in non_blank if isinstance(v, str) and parse_numeric_text(v) is None)
            if n_text * 2 >= len(non_blank):
                header = i
                break
    else:
        header = next(i for i, c in enumerate(counts) if c > 0)
    title_parts = [truncate(str(v).strip(), 80) for row in window[:header] for v in row if not is_blank(v)]
    return header, truncate(" / ".join(title_parts), 200)


def build_header_names(
    header_values: List[Any],
    header_row_number: int,
    merged: List[Tuple[int, int, int, int]],
) -> List[str]:
    """Header text per column; merged header cells are forward-filled from the
    top-left cell, blank headers get their Excel column letter, duplicates get ' (n)'."""
    names: List[str] = []
    seen: Dict[str, int] = {}
    for j, value in enumerate(header_values):
        text = "" if is_blank(value) else truncate(str(value).strip(), 80)
        if not text:
            col = j + 1
            for c1, r1, c2, r2 in merged:
                if r1 <= header_row_number <= r2 and c1 <= col <= c2 and not (c1 == col and r1 == header_row_number):
                    top_left = header_values[c1 - 1] if c1 - 1 < len(header_values) else None
                    if not is_blank(top_left):
                        text = truncate(str(top_left).strip(), 80)
                    break
        if not text:
            text = f"({get_column_letter(j + 1)})"
        base = text
        if base in seen:
            seen[base] += 1
            text = f"{base} ({seen[base]})"
        else:
            seen[base] = 1
        names.append(text)
    return names


# ---------------------------------------------------------------------------
# sheet extraction
# ---------------------------------------------------------------------------
def _row_payload(cells: Tuple[Any, ...]) -> Tuple[List[Any], Set[int], Set[int], Set[int]]:
    """Cell values of one row plus the column indexes that are percent-formatted,
    Excel error cells (value replaced by None) and text placeholders (value replaced by None)."""
    values: List[Any] = []
    pct: Set[int] = set()
    errors: Set[int] = set()
    placeholders: Set[int] = set()
    for j, cell in enumerate(cells):
        value = cell.value
        if getattr(cell, "data_type", None) == "e" or is_error_text(value):
            errors.add(j)
            value = None
        elif is_placeholder(value):
            placeholders.add(j)
            value = None
        values.append(value)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            fmt = getattr(cell, "number_format", None)
            if fmt and "%" in fmt:
                pct.add(j)
    return values, pct, errors, placeholders


def _all_blank(values: List[Any]) -> bool:
    return all(is_blank(v) for v in values)


class _ChartState:
    """Streaming aggregators for the charts chosen from the sample window."""

    def __init__(self, cat: Optional[int], num: Optional[int], date: Optional[int]) -> None:
        self.cat, self.num, self.date = cat, num, date
        self.cat_sum: Dict[str, float] = {}
        self.month_sum: Dict[str, float] = {}
        self.month_count: Dict[str, int] = {}
        self.day_sum: Dict[str, float] = {}
        self.day_count: Dict[str, int] = {}
        self.days_valid = True

    def feed(self, values: List[Any]) -> None:
        get = lambda j: values[j] if j is not None and j < len(values) else None  # noqa: E731
        num_value = numeric_of(get(self.num)) if self.num is not None else None
        if self.cat is not None and num_value is not None:
            cat_value = get(self.cat)
            if not is_blank(cat_value):
                key = truncate(str(cat_value).strip(), 80)
                self.cat_sum[key] = self.cat_sum.get(key, 0.0) + num_value
        if self.date is not None:
            date_value = get(self.date)
            mk = month_key(date_value)
            if mk is not None:
                self.month_count[mk] = self.month_count.get(mk, 0) + 1
                if num_value is not None:
                    self.month_sum[mk] = self.month_sum.get(mk, 0.0) + num_value
                if self.days_valid:
                    dk = day_key(date_value)
                    if dk not in self.day_count and len(self.day_count) >= MAX_DAY_BUCKETS:
                        self.days_valid = False
                        self.day_count.clear()
                        self.day_sum.clear()
                    else:
                        self.day_count[dk] = self.day_count.get(dk, 0) + 1
                        if num_value is not None:
                            self.day_sum[dk] = self.day_sum.get(dk, 0.0) + num_value


def _choose_chart_columns(accs: List[ColumnAcc], max_categories: int, skip_sensitive: bool) -> _ChartState:
    cat = num = date = None
    for acc in accs:
        kind = acc.inferred_type()
        distinct = len(acc.counter)
        if (
            cat is None
            and kind == "text"
            and not acc.distinct_overflow
            and 2 <= distinct <= max(max_categories * 4, 50)
            and distinct * 2 <= acc.n_nonempty
            and not (skip_sensitive and acc.looks_sensitive())
        ):
            cat = acc.index
        elif num is None and kind == "number":
            num = acc.index
        elif date is None and kind in ("date", "datetime"):
            date = acc.index
    return _ChartState(cat, num, date)


def _finish_charts(state: Optional[_ChartState], accs: List[ColumnAcc], kinds: Dict[int, str], names: Dict[int, str], max_categories: int) -> List[Dict[str, Any]]:
    charts: List[Dict[str, Any]] = []
    if state is None:
        return charts
    by_index = {acc.index: acc for acc in accs}
    cat_ok = state.cat is not None and kinds.get(state.cat) == "text"
    num_ok = state.num is not None and kinds.get(state.num) == "number"
    date_ok = state.date is not None and kinds.get(state.date) in ("date", "datetime")
    if cat_ok:
        counter = by_index[state.cat].counter
        labels, values = bucket_categories({k: float(v) for k, v in counter.items()}, max_categories)
        charts.append(
            {
                "kind": "pie",
                "title": f"{names[state.cat]} 분포 (건수)",
                "category": names[state.cat],
                "measure": "건수",
                "labels": labels,
                "values": values,
            }
        )
        if num_ok and state.cat_sum:
            labels, values = bucket_categories(state.cat_sum, max_categories)
            charts.append(
                {
                    "kind": "bar",
                    "title": f"{names[state.cat]}별 {names[state.num]} 합계",
                    "category": names[state.cat],
                    "measure": names[state.num],
                    "labels": labels,
                    "values": values,
                }
            )
    if date_ok and state.month_count:
        use_days = state.days_valid and 0 < len(state.day_count) <= MAX_LINE_POINTS
        source_count = state.day_count if use_days else state.month_count
        source_sum = state.day_sum if use_days else state.month_sum
        keys = sorted(source_count)
        truncated = len(keys) > MAX_LINE_POINTS
        keys = keys[-MAX_LINE_POINTS:]
        if num_ok:
            values = [round(source_sum.get(k, 0.0), 6) for k in keys]
            measure = f"{names[state.num]} 합계"
        else:
            values = [source_count[k] for k in keys]
            measure = "건수"
        values = [int(v) if isinstance(v, float) and v.is_integer() else v for v in values]
        charts.append(
            {
                "kind": "line",
                "title": f"{names[state.date]}별 {measure}" + (" (일 단위)" if use_days else " (월 단위)"),
                "category": names[state.date],
                "measure": measure,
                "labels": keys,
                "values": values,
                "truncated": truncated,
            }
        )
    return charts


def _count_formulas(ws_formula: Any, first_row: int, last_row: int, kept: List[int], buffered: Dict[int, List[Any]]) -> Tuple[int, int]:
    """Count formula cells in the published window, and how many of them have no cached value."""
    n_formula = 0
    n_without = 0
    if ws_formula is None or last_row < first_row:
        return 0, 0
    row_number = first_row - 1
    for cells in ws_formula.iter_rows(min_row=first_row, max_row=last_row):
        row_number += 1
        for j in kept:
            if j >= len(cells):
                continue
            if getattr(cells[j], "data_type", None) == "f":
                n_formula += 1
                values = buffered.get(row_number)
                if values is None or j >= len(values) or is_blank(values[j]):
                    n_without += 1
    return n_formula, n_without


def _empty_sheet(sheet_name: str) -> Dict[str, Any]:
    return {
        "name": sheet_name,
        "state": "empty",
        "title": "",
        "header_row": None,
        "rows_total": 0,
        "rows_scanned": 0,
        "rows_published": 0,
        "scan_truncated": False,
        "columns_total": 0,
        "columns_published": 0,
        "columns_truncated": 0,
        "columns_excluded": [],
        "cells_truncated": 0,
        "formula_cells": 0,
        "formula_cells_without_cached_value": 0,
        "columns": [],
        "rows": [],
        "charts": [],
        "notes": ["시트에 데이터가 없습니다 (빈 시트)."],
    }


def extract_sheet(ws: Any, ws_formula: Any, sheet_name: str, cfg: Dict[str, Any], merged: List[Tuple[int, int, int, int]]) -> Dict[str, Any]:
    header_scan = cfg["header_scan_rows"]
    max_rows = cfg["max_rows"]
    max_scan = cfg["max_scan_rows"]
    max_cell_text = cfg["max_cell_text"]
    publish_rows = cfg["mode"] == "rows"
    auto_exclude = cfg["auto_exclude_sensitive"]
    allow = cfg["allow_columns"]

    rows_iter = ws.iter_rows()
    window: List[Tuple[int, List[Any], Set[int], Set[int], Set[int]]] = []
    physical = 0
    for cells in rows_iter:
        physical += 1
        values, pct, errors, placeholders = _row_payload(cells)
        if _all_blank(values) and not errors:
            if physical >= MAX_HEADER_SEARCH_ROWS and not window:
                break
            continue
        window.append((physical, values, pct, errors, placeholders))
        if len(window) >= header_scan:
            break
    if not window or all(_all_blank(w[1]) for w in window):
        return _empty_sheet(sheet_name)

    header_pos, title = detect_header([w[1] for w in window])
    header_row_number = window[header_pos][0]
    header_values = list(window[header_pos][1])
    width = max(len(w[1]) for w in window)
    header_values += [None] * (width - len(header_values))
    names = build_header_names(header_values, header_row_number, merged)

    # column selection by config and by sensitive header names, then hard cap for the scan
    excluded: List[Dict[str, str]] = []
    kept: List[int] = []
    for j, name in enumerate(names):
        if cfg["include_columns"] and not name_matches(name, cfg["include_columns"]):
            excluded.append({"name": name, "reason": "config"})
            continue
        if name_matches(name, cfg["exclude_columns"]):
            excluded.append({"name": name, "reason": "config"})
            continue
        if auto_exclude and not name_matches(name, allow) and sensitive_name(name, cfg["sensitive_patterns"]):
            excluded.append({"name": name, "reason": "sensitive-name"})
            continue
        kept.append(j)
    kept = kept[:HARD_COLUMN_CAP]
    accs = [ColumnAcc(names[j], get_column_letter(j + 1), j) for j in kept]

    buffer: List[Tuple[int, List[Any]]] = []
    state: Dict[str, Any] = {"total": 0, "scanned": 0, "truncated": False}
    chart: Optional[_ChartState] = None

    def choose_charts() -> None:
        nonlocal chart
        chart = _choose_chart_columns(accs, cfg["max_chart_categories"], auto_exclude)
        for _, values in buffer:
            chart.feed(values)

    def handle(row_number: int, values: List[Any], pct: Set[int], errors: Set[int], placeholders: Set[int]) -> None:
        nonlocal chart
        if _all_blank(values) and not errors:
            return
        state["total"] += 1
        if state["scanned"] >= max_scan:
            state["truncated"] = True
            return
        state["scanned"] += 1
        for acc in accs:
            j = acc.index
            if j in errors:
                acc.note_error()
                continue
            if j in placeholders:
                acc.note_placeholder()
                continue
            value = values[j] if j < len(values) else None
            acc.feed(value, "0%" if j in pct else "General")
        if len(buffer) < max_rows:
            buffer.append((row_number, values))
            if len(buffer) == max_rows and chart is None:
                choose_charts()
        elif chart is None:
            choose_charts()
            chart.feed(values)
        else:
            chart.feed(values)

    for row_number, values, pct, errors, placeholders in window[header_pos + 1 :]:
        handle(row_number, values, pct, errors, placeholders)
    for cells in rows_iter:
        physical += 1
        values, pct, errors, placeholders = _row_payload(cells)
        handle(physical, values, pct, errors, placeholders)
    if chart is None:
        choose_charts()

    kinds_all = {acc.index: acc.inferred_type() for acc in accs}

    # drop unnamed, entirely empty columns (typical for over-wide sheet dimensions)
    live = [acc for acc in accs if not (acc.n_nonempty == 0 and acc.n_error == 0 and acc.name.startswith("(") and acc.name.endswith(")"))]
    # drop columns whose VALUES look like personal data (decided after the scan)
    if auto_exclude:
        remaining: List[ColumnAcc] = []
        for acc in live:
            if acc.looks_sensitive() and not name_matches(acc.name, allow):
                excluded.append({"name": acc.name, "reason": "sensitive-values"})
            else:
                remaining.append(acc)
        live = remaining
    columns_total = len(live) + len(excluded)
    columns_truncated = max(0, len(live) - cfg["max_columns"])
    live = live[: cfg["max_columns"]]
    live_indexes = [acc.index for acc in live]
    kinds = {acc.index: kinds_all[acc.index] for acc in live}
    name_by_index = {acc.index: acc.name for acc in live}
    top_values = publish_rows and cfg["publish_top_values"]

    columns = [
        {
            "name": acc.name,
            "excel_column": acc.letter,
            "type": kinds[acc.index],
            "non_empty": acc.n_nonempty,
            "errors": acc.n_error,
            "stats": acc.stats(kinds[acc.index], top_values=top_values),
        }
        for acc in live
    ]
    rows_out: List[List[Any]] = []
    cells_truncated = 0
    if publish_rows:
        for _, values in buffer:
            row: List[Any] = []
            for j in live_indexes:
                v = convert_value(values[j] if j < len(values) else None, kinds[j])
                if isinstance(v, str) and len(v) > max_cell_text:
                    v = v[: max_cell_text - 1] + "…"
                    cells_truncated += 1
                row.append(v)
            rows_out.append(row)

    last_buffered = buffer[-1][0] if buffer else header_row_number
    n_formula, n_without = _count_formulas(
        ws_formula, header_row_number + 1, last_buffered, live_indexes, {r: v for r, v in buffer}
    )
    charts = _finish_charts(chart, live, kinds, name_by_index, cfg["max_chart_categories"]) if cfg["charts"] else []

    notes: List[str] = []
    if state["truncated"]:
        notes.append(f"집계는 처음 {max_scan:,}행만 사용했습니다 (총 {state['total']:,}행).")
    if publish_rows and state["total"] > len(buffer):
        notes.append(f"행 데이터는 처음 {len(buffer):,}행만 공개합니다 (총 {state['total']:,}행).")
    if not publish_rows:
        notes.append("집계 전용 모드: 행 데이터는 공개하지 않습니다 (열 이름·통계·차트 범주만 공개).")
    if columns_truncated:
        notes.append(f"열은 처음 {cfg['max_columns']}개만 공개합니다 ({columns_truncated}개 열 생략).")
    for reason, label in EXCLUSION_REASONS.items():
        names_for_reason = [e["name"] for e in excluded if e["reason"] == reason]
        if names_for_reason:
            notes.append(f"{label}: " + ", ".join(names_for_reason[:20]) + (" …" if len(names_for_reason) > 20 else ""))
    n_errors = sum(acc.n_error for acc in live)
    if n_errors:
        notes.append(f"오류 셀 {n_errors:,}개(#DIV/0!, #N/A 등)는 빈 값으로 처리했습니다.")
    n_placeholders = sum(acc.n_placeholder for acc in live)
    if n_placeholders:
        notes.append(f"'-', 'N/A' 같은 자리표시 값 {n_placeholders:,}개는 빈 값으로 처리했습니다.")
    if cells_truncated:
        notes.append(f"{max_cell_text}자를 넘는 텍스트 셀 {cells_truncated:,}개는 잘라서 공개합니다.")
    if n_formula:
        notes.append(
            f"수식 셀 {n_formula:,}개 중 캐시된 값이 없는 셀 {n_without:,}개 "
            "(Excel에서 저장한 파일만 계산 결과가 들어 있습니다)."
        )
    coerced = [c["name"] for c in columns if c["stats"].get("coerced_from_text")]
    if coerced:
        notes.append("텍스트로 저장된 숫자/날짜를 변환한 열: " + ", ".join(coerced[:20]))
    for c in charts:
        if c.get("truncated"):
            notes.append(f"시계열 차트는 마지막 {MAX_LINE_POINTS}개 구간만 표시합니다.")

    return {
        "name": sheet_name,
        "state": "ok",
        "title": title,
        "header_row": header_row_number,
        "rows_total": state["total"],
        "rows_scanned": state["scanned"],
        "rows_published": len(rows_out),
        "scan_truncated": state["truncated"],
        "columns_total": columns_total,
        "columns_published": len(columns),
        "columns_truncated": columns_truncated,
        "columns_excluded": excluded,
        "cells_truncated": cells_truncated,
        "formula_cells": n_formula,
        "formula_cells_without_cached_value": n_without,
        "columns": columns,
        "rows": rows_out,
        "charts": charts,
        "notes": notes,
    }


def extract_workbook(path: Path, rel_path: str, cfg: Dict[str, Any]) -> Dict[str, Any]:
    """Extract every visible, selected worksheet of one workbook. Raises BuildError on failure."""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # BadZipFile, InvalidFileException, KeyError, ...
        raise BuildError(f"{rel_path}: cannot open workbook ({type(exc).__name__}: {exc})") from exc
    try:
        wb_formula = load_workbook(path, read_only=True, data_only=False)
    except Exception:
        wb_formula = None
    merged_all = header_merged_ranges(path, cfg["header_scan_rows"] + 5)
    sheets: List[Dict[str, Any]] = []
    n_hidden = 0
    n_excluded = 0
    try:
        for ws in wb.worksheets:
            name = ws.title
            if getattr(ws, "sheet_state", "visible") != "visible":
                n_hidden += 1  # hidden / veryHidden sheets: neither content nor name is published
                continue
            if cfg["include_sheets"] and not name_matches(name, cfg["include_sheets"]):
                n_excluded += 1
                continue
            if name_matches(name, cfg["exclude_sheets"]):
                n_excluded += 1
                continue
            ws_formula = None
            if wb_formula is not None and name in wb_formula.sheetnames:
                ws_formula = wb_formula[name]
            try:
                sheets.append(extract_sheet(ws, ws_formula, name, cfg, merged_all.get(name, [])))
            except BuildError:
                raise
            except Exception as exc:
                raise BuildError(f"{rel_path} / sheet '{name}': cannot read sheet ({type(exc).__name__}: {exc})") from exc
    finally:
        wb.close()
        if wb_formula is not None:
            wb_formula.close()
    notes: List[str] = []
    if n_hidden:
        notes.append(f"숨김 시트 {n_hidden}개는 공개하지 않습니다.")
    if n_excluded:
        notes.append(f"설정으로 제외한 시트 {n_excluded}개는 공개하지 않습니다.")
    return {
        "path": rel_path,
        "fingerprint": fingerprint([[s["name"], s["columns"], s["rows"], s["rows_total"]] for s in sheets]),
        "sheets": sheets,
        "sheets_hidden": n_hidden,
        "sheets_excluded": n_excluded,
        "notes": notes,
    }
