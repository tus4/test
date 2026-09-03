"""CLI: regenerate the clearly-labelled synthetic sample site (docs/ + README block).

    python -m dashboard_builder.sample --repo-root .

The sample workbook is generated deterministically (no randomness, no clock),
so the committed docs/ can be reproduced byte-for-byte; CI checks exactly that
while docs/ still holds the sample.
"""

from __future__ import annotations

import argparse
import datetime as dt
import sys
import tempfile
from pathlib import Path
from typing import Optional

from openpyxl import Workbook

from .build import build
from .common import BuildError
from .publish import apply_staging

SAMPLE_PAGES_URL = "https://tus4.github.io/test/"
SAMPLE_CONFIG = '{\n  "title": "데이터 대시보드 (샘플)",\n  "preview_rows": 8\n}\n'


def make_sample_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "월별 매출"
    ws["A1"] = "샘플 데이터 (SAMPLE) - 지역별 월 매출 현황"
    ws.merge_cells("A1:G1")
    ws["A2"] = "이 워크북은 자동 생성된 예시이며 실제 데이터가 아닙니다."
    ws.append([])
    ws.append(["일자", "지역", "제품", "수량", "단가", "매출액", "목표 달성률"])
    regions = ["서울", "부산", "대구", "광주"]
    products = ["노트북", "모니터", "키보드"]
    row = 5
    for month in range(1, 13):
        for r_i, region in enumerate(regions):
            for p_i, product in enumerate(products):
                qty = 10 + ((month * 7 + r_i * 3 + p_i * 5) % 23)
                unit = [1200000, 350000, 45000][p_i]
                ws.cell(row=row, column=1, value=dt.datetime(2024, month, 1 + (r_i * 6 + p_i) % 27))
                ws.cell(row=row, column=2, value=region)
                ws.cell(row=row, column=3, value=product)
                ws.cell(row=row, column=4, value=qty)
                ws.cell(row=row, column=5, value=unit)
                ws.cell(row=row, column=6, value=qty * unit)
                c = ws.cell(row=row, column=7, value=round(0.6 + ((month + r_i + p_i) % 9) * 0.06, 2))
                c.number_format = "0%"
                row += 1

    ws2 = wb.create_sheet("거래처")
    # "연락처" is a personal-data column: the build excludes it automatically (auto_exclude_sensitive).
    ws2.append(["거래처명", "구분", "담당자", "연락처", "월 거래액", "등록일"])
    kinds = ["도매", "소매", "온라인"]
    for i in range(1, 25):
        ws2.append(
            [
                f"거래처 {i:02d} (SAMPLE)",
                kinds[i % 3],
                "-" if i == 7 else f"담당자{i % 5 + 1}",  # "-" placeholder -> published as empty
                f"010-{1000 + i * 37:04d}-{2000 + i * 53:04d}",
                f"{(i * 137) % 900 + 100},000",  # numbers stored as text
                dt.datetime(2023, (i % 12) + 1, (i % 28) + 1),
            ]
        )
    error_cell = ws2.cell(row=5, column=5, value="#N/A")  # an Excel error cell -> published as empty
    error_cell.data_type = "e"
    wb.create_sheet("메모")  # empty sheet
    wb.save(path)


def regenerate(repo_root: Path) -> None:
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        source = tmp_path / "data1"
        source.mkdir()
        make_sample_workbook(source / "sample-data.xlsx")
        (source / "dashboard.config.json").write_text(SAMPLE_CONFIG, encoding="utf-8")
        staging = tmp_path / "staging"
        build(source, staging, None, SAMPLE_PAGES_URL, sample=True)
        apply_staging(staging, repo_root)


def main(argv: Optional[list] = None) -> int:
    parser = argparse.ArgumentParser(description="Regenerate the synthetic sample dashboard.")
    parser.add_argument("--repo-root", default=".", help="root of the public repository checkout")
    args = parser.parse_args(argv)
    try:
        regenerate(Path(args.repo_root).resolve())
    except BuildError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    print("Sample dashboard regenerated in docs/ and README.md.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
