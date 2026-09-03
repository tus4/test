"""Shared fixtures: synthetic workbooks are generated here, never checked in."""

from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def make_hard_workbook(path: Path) -> None:
    """Title rows, merged cells, sparse group-header row, dates, formulas, numbers-as-text,
    percent format, awkward characters, an empty sheet, a one-cell sheet and a hidden sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "매출현황"
    ws["A1"] = "2024년 지점별 매출 현황 (단위: 원)"
    ws.merge_cells("A1:I1")
    ws["A2"] = "작성: 기획팀"
    ws["A4"] = "기본 정보"
    ws.merge_cells("A4:C4")
    ws["D4"] = "판매"
    ws.merge_cells("D4:F4")
    ws["G4"] = "기타"
    ws.merge_cells("G4:I4")
    ws.append(["일자", "지점", "상품", "수량", "단가", "금액(수식)", "텍스트숫자", "달성률", "비고"])
    branches = ["서울", "부산", "대구", "광주"]
    products = ["A상품", "B상품", "C상품", "D상품", "E상품", "F상품"]
    for i in range(40):
        r = 6 + i
        ws.cell(r, 1, dt.datetime(2024, 1 + i % 12, 1 + (i * 3) % 28))
        ws.cell(r, 2, branches[i % 4])
        ws.cell(r, 3, products[i % 6])
        ws.cell(r, 4, 5 + (i * 7) % 20)
        ws.cell(r, 5, 1000 * (1 + i % 5))
        ws.cell(r, 6, f"=D{r}*E{r}")
        ws.cell(r, 7, f"{(i + 1) * 1234:,}")
        c = ws.cell(r, 8, round(0.5 + (i % 10) * 0.05, 2))
        c.number_format = "0%"
        ws.cell(r, 9, ["정상", "지연|보류", "재고\n부족", None][i % 4])
    wb.create_sheet("빈시트")
    memo = wb.create_sheet("메모")
    memo["B2"] = "메모 한 줄"
    hidden = wb.create_sheet("숨김")
    hidden.append(["비밀", "값"])
    hidden.append(["x", 1])
    hidden.sheet_state = "hidden"
    wb.save(path)


def make_wide_long_workbook(path: Path, rows: int = 1500, extra_cols: int = 55) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "재고"
    ws.append(["품목코드", "품목명", "창고", "수량", "단가"] + [f"속성{i}" for i in range(1, extra_cols + 1)])
    for i in range(rows):
        ws.append(
            [f"P{i:05d}", f"품목 {i}", ["A창고", "B창고", "C창고"][i % 3], (i * 13) % 500, 100 + (i % 50) * 10]
            + [(i * j) % 97 for j in range(1, extra_cols + 1)]
        )
    wb.save(path)


def make_simple_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["id", "value", "when"])
    ws.append([1, 2.5, dt.datetime(2024, 5, 1, 9, 30)])
    ws.append([2, 3.5, dt.datetime(2024, 5, 2)])
    wb.save(path)


@pytest.fixture
def data_dir(tmp_path: Path) -> Path:
    """A simulated data1 checkout with three workbooks, a lock file and a .git dir."""
    src = tmp_path / "data1"
    (src / "2024").mkdir(parents=True)
    (src / ".git").mkdir()
    (src / ".git" / "HEAD").write_text("ref: refs/heads/main\n")
    make_hard_workbook(src / "hard.xlsx")
    make_wide_long_workbook(src / "2024" / "재고.xlsx")
    make_simple_workbook(src / "macro.xlsm")
    (src / "~$hard.xlsx").write_bytes(b"\x00lock")
    (src / "notes.txt").write_text("not a workbook")
    return src
