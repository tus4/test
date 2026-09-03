import datetime as dt
from pathlib import Path

import pytest
from openpyxl import Workbook

from dashboard_builder.analyze import ColumnAcc, parse_numeric_text
from dashboard_builder.common import BuildError
from dashboard_builder.config import validate_config
from dashboard_builder.excel import (
    build_header_names,
    detect_header,
    discover_workbooks,
    extract_workbook,
    header_merged_ranges,
    path_matches,
)

from conftest import make_hard_workbook


def cfg(**overrides):
    return validate_config(overrides)


def test_discovery_skips_lock_files_hidden_dirs_and_non_workbooks(data_dir: Path) -> None:
    found = discover_workbooks(data_dir, cfg())
    assert [rel for _, rel in found] == ["2024/재고.xlsx", "hard.xlsx", "macro.xlsm"]


def test_discovery_respects_include_exclude(data_dir: Path) -> None:
    assert [rel for _, rel in discover_workbooks(data_dir, cfg(exclude_files=["2024/**"]))] == ["hard.xlsx", "macro.xlsm"]
    assert [rel for _, rel in discover_workbooks(data_dir, cfg(include_files=["*.xlsm"]))] == ["macro.xlsm"]


def test_path_matches_root_and_nested() -> None:
    assert path_matches("a.xlsx", ["**/*.xlsx"])
    assert path_matches("x/y/a.xlsx", ["**/*.xlsx"])
    assert not path_matches("x/a.xlsm", ["**/*.xlsx"])
    assert path_matches("backup/a.xlsx", ["backup/**"])


def test_path_matches_is_case_insensitive() -> None:
    # An upper-case extension must not fall out of the DEFAULT include_files patterns.
    assert path_matches("logs/센서.XLSX", ["**/*.xlsx", "**/*.xlsm"])
    assert path_matches("Report.Xlsm", ["**/*.xlsx", "**/*.xlsm"])
    assert path_matches("Backup/old.xlsx", ["backup/**"])


def test_discovery_keeps_upper_case_extension_with_default_config(tmp_path: Path) -> None:
    from dashboard_builder.excel import discover_workbooks_with_excluded

    (tmp_path / "logs").mkdir()
    wb = Workbook()
    wb.active.append(["a", "b"])
    wb.active.append([1, 2])
    wb.save(tmp_path / "logs" / "UPPER.XLSX")
    wb.save(tmp_path / "lower.xlsx")
    found, excluded = discover_workbooks_with_excluded(tmp_path, cfg())
    assert [rel for _, rel in found] == ["logs/UPPER.XLSX", "lower.xlsx"]
    assert excluded == 0  # nothing is "excluded by config" when there is no config
    found, excluded = discover_workbooks_with_excluded(tmp_path, cfg(exclude_files=["LOGS/**"]))
    assert [rel for _, rel in found] == ["lower.xlsx"] and excluded == 1


def test_parse_numeric_text() -> None:
    assert parse_numeric_text("1,234.5") == (1234.5, False)
    assert parse_numeric_text(" 15% ") == (0.15, True)
    assert parse_numeric_text("-42") == (-42.0, False)
    assert parse_numeric_text("abc") is None
    assert parse_numeric_text("1,23") is None


def test_detect_header_prefers_dense_text_row() -> None:
    window = [["제목", None, None], ["소제목", None, None], ["a", "b", "c"], [1, 2, 3]]
    idx, title = detect_header(window)
    assert idx == 2
    assert title == "제목 / 소제목"


def test_detect_header_single_column_sheet() -> None:
    idx, title = detect_header([["only"], ["x"]])
    assert idx == 0 and title == ""


def test_header_names_merged_blank_and_duplicates() -> None:
    names = build_header_names(["판매", None, None, "수량", "수량"], 5, [(1, 5, 2, 5)])
    assert names == ["판매", "판매 (2)", "(C)", "수량", "수량 (2)"]


def test_merged_ranges_are_read_from_xml(tmp_path: Path) -> None:
    path = tmp_path / "m.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "t"
    ws.merge_cells("A1:C1")
    ws.merge_cells("A30:B31")  # outside the header window
    wb.save(path)
    assert header_merged_ranges(path, 20) == {"S": [(1, 1, 3, 1)]}


def test_column_acc_types() -> None:
    num = ColumnAcc("n", "A", 0)
    for v in ["1,000", "2,500", "10"]:
        num.feed(v)
    assert num.inferred_type() == "number"
    assert num.stats("number")["coerced_from_text"] == 3
    pct = ColumnAcc("p", "B", 1)
    pct.feed(0.5, "0%")
    pct.feed(0.25, "0.0%")
    assert pct.inferred_type() == "percent"
    mixed = ColumnAcc("m", "C", 2)
    mixed.feed(1)
    mixed.feed("abc")
    assert mixed.inferred_type() == "text"
    dates = ColumnAcc("d", "D", 3)
    dates.feed(dt.datetime(2024, 1, 1))
    dates.feed(dt.datetime(2024, 1, 2, 10, 0))
    assert dates.inferred_type() == "datetime"
    empty = ColumnAcc("e", "E", 4)
    empty.feed(None)
    empty.feed("   ")
    assert empty.inferred_type() == "empty"


def test_hard_workbook_extraction(tmp_path: Path) -> None:
    path = tmp_path / "hard.xlsx"
    make_hard_workbook(path)
    file = extract_workbook(path, "hard.xlsx", cfg())
    assert file["notes"][0] == "숨김 시트 1개는 공개하지 않습니다." and file["sheets_hidden"] == 1
    by_name = {s["name"]: s for s in file["sheets"]}
    assert set(by_name) == {"매출현황", "빈시트", "메모"}  # hidden sheet is not published

    s = by_name["매출현황"]
    assert s["header_row"] == 5
    assert s["title"].startswith("2024년 지점별 매출 현황")
    assert [c["name"] for c in s["columns"]] == ["일자", "지점", "상품", "수량", "단가", "금액(수식)", "텍스트숫자", "달성률", "비고"]
    types = {c["name"]: c["type"] for c in s["columns"]}
    assert types["일자"] == "date"
    assert types["지점"] == "text"
    assert types["수량"] == "number"
    assert types["금액(수식)"] == "empty"  # formula without cached value (openpyxl-written file)
    assert types["텍스트숫자"] == "number"
    assert types["달성률"] == "percent"
    assert types["비고"] == "text"
    assert s["formula_cells"] == 40 and s["formula_cells_without_cached_value"] == 40
    assert s["rows_total"] == 40 and s["rows_published"] == 40
    assert s["rows"][0][0] == "2024-01-01" and s["rows"][0][6] == 1234 and s["rows"][0][7] == 0.5
    assert s["rows"][1][8] == "지연|보류" and s["rows"][2][8] == "재고\n부족"
    stats = {c["name"]: c["stats"] for c in s["columns"]}
    assert stats["수량"]["sum"] == sum(5 + (i * 7) % 20 for i in range(40))
    assert stats["일자"]["min"] == "2024-01-01"
    kinds = [c["kind"] for c in s["charts"]]
    assert kinds == ["pie", "bar", "line"]
    pie = s["charts"][0]
    assert sorted(pie["labels"]) == ["광주", "대구", "부산", "서울"] and pie["values"] == [10, 10, 10, 10]

    assert by_name["빈시트"]["state"] == "empty"
    memo = by_name["메모"]
    assert memo["rows_total"] == 0 and memo["columns_published"] == 1


def test_caps_and_notes(data_dir: Path) -> None:
    file = extract_workbook(data_dir / "2024" / "재고.xlsx", "2024/재고.xlsx", cfg(max_rows=100, max_columns=10, max_scan_rows=1200))
    s = file["sheets"][0]
    assert s["rows_total"] == 1500 and s["rows_scanned"] == 1200 and s["scan_truncated"] is True
    assert s["rows_published"] == 100 and len(s["rows"]) == 100
    assert s["columns_published"] == 10 and s["columns_truncated"] == 50 and s["columns_total"] == 60
    assert all(len(r) == 10 for r in s["rows"])
    joined = " ".join(s["notes"])
    assert "1,200" in joined and "100" in joined and "10개" in joined


def test_aggregates_mode_has_no_rows(data_dir: Path) -> None:
    file = extract_workbook(data_dir / "hard.xlsx", "hard.xlsx", cfg(mode="aggregates"))
    s = file["sheets"][0]
    assert s["rows"] == [] and s["rows_published"] == 0
    assert s["charts"] and s["columns"]  # aggregates still present


def test_column_and_sheet_filters(data_dir: Path) -> None:
    c = cfg(exclude_columns=["*수식*", "비고"], exclude_sheets=["메모", "빈*"])
    file = extract_workbook(data_dir / "hard.xlsx", "hard.xlsx", c)
    assert [s["name"] for s in file["sheets"]] == ["매출현황"]
    s = file["sheets"][0]
    assert "금액(수식)" not in [col["name"] for col in s["columns"]]
    assert [e["name"] for e in s["columns_excluded"]] == ["금액(수식)", "비고"]
    assert all(e["reason"] == "config" for e in s["columns_excluded"])
    c2 = cfg(include_columns=["일자", "수량"])
    s2 = extract_workbook(data_dir / "hard.xlsx", "hard.xlsx", c2)["sheets"][0]
    assert [col["name"] for col in s2["columns"]] == ["일자", "수량"]


def test_broken_workbook_raises_with_file_name(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a zip")
    with pytest.raises(BuildError, match="broken.xlsx"):
        extract_workbook(path, "broken.xlsx", cfg())
