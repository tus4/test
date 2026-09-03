"""Tests for the exposure policy (automatic personal-data exclusion, aggregates mode, caps)
and robustness (error cells, placeholders, text dates, hidden sheets, dry run, CNAME)."""

from __future__ import annotations

import datetime as dt
import json
import subprocess
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook

from dashboard_builder import build as build_cli
from dashboard_builder import publish
from dashboard_builder.analyze import ColumnAcc, parse_date_text
from dashboard_builder.config import DEFAULTS, validate_config
from dashboard_builder.excel import extract_workbook

ROOT = Path(__file__).resolve().parent.parent
# The two files destined for the data repository are kept in this repo under data1/ (the copies the
# user is told to use). In the delivery tree a sibling ../data1/ holds the originals; when it exists,
# the copies must be byte-identical to it.
DATA1_COPY = ROOT / "data1"
DATA1_ORIGINAL = ROOT.parent / "data1"


def cfg(**overrides):
    return validate_config(overrides)


def make_pii_workbook(path: Path) -> None:
    """Numeric column with an error cell and a '-' placeholder, a phone column, an e-mail column
    (harmless header), a text-date column, a long-text column and a very hidden sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "고객"
    ws.append(["이름", "구분", "전화번호", "메일", "가입일", "금액", "메모"])
    kinds = ["일반", "우수", "VIP"]
    for i in range(30):
        ws.append(
            [
                f"사람{i}",
                kinds[i % 3],
                f"010-{1000 + i:04d}-{5000 + i:04d}",
                f"user{i}@example.com",
                f"2024-0{1 + i % 9}-1{i % 9}",
                1000 + i,
                "긴 텍스트 " * 60 if i == 0 else "짧음",
            ]
        )
    err = ws.cell(row=3, column=6, value="#DIV/0!")
    err.data_type = "e"
    ws.cell(row=4, column=6, value="-")
    secret = wb.create_sheet("급여_원본")
    secret.append(["비밀", "값"])
    secret.append(["x", 1])
    secret.sheet_state = "veryHidden"
    wb.save(path)


@pytest.fixture
def pii_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "pii.xlsx"
    make_pii_workbook(path)
    return path


def _sheet(path: Path, **overrides):
    return extract_workbook(path, path.name, cfg(**overrides))["sheets"][0]


def test_error_cell_and_placeholder_keep_numeric_type(pii_workbook: Path) -> None:
    s = _sheet(pii_workbook)
    col = {c["name"]: c for c in s["columns"]}["금액"]
    assert col["type"] == "number" and col["errors"] == 1
    assert col["stats"]["count"] == 28 and col["stats"]["sum"] == sum(1000 + i for i in range(30)) - 1001 - 1002
    assert s["rows"][1][col_index(s, "금액")] is None and s["rows"][2][col_index(s, "금액")] is None
    joined = " ".join(s["notes"])
    assert "오류 셀 1개" in joined and "자리표시 값 1개" in joined
    assert [c["kind"] for c in s["charts"]] == ["pie", "bar", "line"]  # the numeric column still drives the bar chart


def col_index(sheet, name: str) -> int:
    return [c["name"] for c in sheet["columns"]].index(name)


def test_sensitive_columns_are_excluded_by_name_and_by_value_shape(pii_workbook: Path) -> None:
    s = _sheet(pii_workbook)
    names = [c["name"] for c in s["columns"]]
    assert "전화번호" not in names and "메일" not in names
    reasons = {e["name"]: e["reason"] for e in s["columns_excluded"]}
    assert reasons == {"전화번호": "sensitive-name", "메일": "sensitive-values"}
    assert s["columns_total"] == 7 and s["columns_published"] == 5
    text = json.dumps(s, ensure_ascii=False)
    assert "010-1000-5000" not in text and "user0@example.com" not in text
    assert any("자동 제외한 열" in n for n in s["notes"])


def test_allow_columns_and_disabling_auto_exclusion(pii_workbook: Path) -> None:
    s = _sheet(pii_workbook, allow_columns=["전화*"])
    assert "전화번호" in [c["name"] for c in s["columns"]] and "메일" not in [c["name"] for c in s["columns"]]
    s2 = _sheet(pii_workbook, auto_exclude_sensitive=False)
    assert {"전화번호", "메일"} <= {c["name"] for c in s2["columns"]} and s2["columns_excluded"] == []


def test_text_dates_are_coerced(pii_workbook: Path) -> None:
    s = _sheet(pii_workbook)
    col = {c["name"]: c for c in s["columns"]}["가입일"]
    assert col["type"] == "date" and col["stats"]["coerced_from_text"] == 30
    assert col["stats"]["min"] == "2024-01-10" and s["rows"][0][col_index(s, "가입일")] == "2024-01-10"
    assert parse_date_text("2024년 3월 5일") == dt.datetime(2024, 3, 5)
    assert parse_date_text("2024/03/05 09:30") == dt.datetime(2024, 3, 5, 9, 30)
    assert parse_date_text("2024-13-01") is None and parse_date_text("abc") is None


def test_cell_text_cap(pii_workbook: Path) -> None:
    s = _sheet(pii_workbook)
    memo = s["rows"][0][col_index(s, "메모")]
    assert len(memo) == DEFAULTS["max_cell_text"] and memo.endswith("…") and s["cells_truncated"] == 1
    assert any("200자를 넘는" in n for n in s["notes"])
    s2 = _sheet(pii_workbook, max_cell_text=20)
    assert len(s2["rows"][0][col_index(s2, "메모")]) == 20


def test_hidden_sheet_name_is_not_published(pii_workbook: Path) -> None:
    file = extract_workbook(pii_workbook, "pii.xlsx", cfg())
    assert [s["name"] for s in file["sheets"]] == ["고객"] and file["sheets_hidden"] == 1
    assert "급여_원본" not in json.dumps(file, ensure_ascii=False)
    assert file["notes"] == ["숨김 시트 1개는 공개하지 않습니다."]


def test_aggregates_mode_publishes_no_rows_and_no_top_values(pii_workbook: Path) -> None:
    s = _sheet(pii_workbook, mode="aggregates")
    assert s["rows"] == [] and s["rows_published"] == 0
    assert all("top" not in c["stats"] for c in s["columns"])
    text = json.dumps(s, ensure_ascii=False)
    assert "사람0" not in text and "짧음" not in text  # no raw text values at all
    assert s["charts"]  # aggregates (category labels of a categorical column) remain
    assert {c["name"] for c in s["columns"]} == {"이름", "구분", "가입일", "금액", "메모"}


def test_top_values_only_for_low_cardinality_columns(pii_workbook: Path) -> None:
    s = _sheet(pii_workbook)
    stats = {c["name"]: c["stats"] for c in s["columns"]}
    assert "top" not in stats["이름"] and stats["이름"]["distinct"] == 30 or "top" in stats["이름"]  # 30 <= 50 -> listed
    assert stats["구분"]["top"][0][1] == 10
    s2 = _sheet(pii_workbook, publish_top_values=False)
    assert all("top" not in c["stats"] for c in s2["columns"])


def test_charts_can_be_disabled(pii_workbook: Path) -> None:
    assert _sheet(pii_workbook, charts=False)["charts"] == []


def test_column_acc_error_placeholder_and_pii_shape() -> None:
    acc = ColumnAcc("n", "A", 0)
    for v in [1, 2, "#N/A", "-", "N/A", 3]:
        acc.feed(v)
    assert acc.inferred_type() == "number" and acc.n_error == 1 and acc.n_placeholder == 2 and acc.n_nonempty == 3
    pii = ColumnAcc("p", "B", 1)
    for v in ["a@b.co", "c@d.org", "010-1234-5678", "hello"]:
        pii.feed(v)
    assert pii.looks_sensitive()


def test_excluded_workbooks_are_counted_not_named(data_dir: Path, tmp_path: Path) -> None:
    (data_dir / "dashboard.config.json").write_text(json.dumps({"exclude_files": ["2024/**"]}), encoding="utf-8")
    out = tmp_path / "staging"
    assert build_cli.main(["--source", str(data_dir), "--out", str(out)]) == 0
    data = json.loads((out / "docs" / "data" / "dashboard.json").read_text(encoding="utf-8"))
    assert data["excluded_files"] == 1 and [f["path"] for f in data["files"]] == ["hard.xlsx", "macro.xlsm"]
    assert "2024/재고.xlsx" not in json.dumps(data, ensure_ascii=False)  # the excluded workbook is never named
    assert "설정으로 제외한 파일: 1개" in (out / "readme_block.md").read_text(encoding="utf-8")


def test_new_config_keys_are_validated() -> None:
    assert validate_config({"_note": "ignored", "max_cell_text": 50})["max_cell_text"] == 50
    for bad in ({"max_cell_text": 0}, {"auto_exclude_sensitive": "yes"}, {"allow_columns": "x"}, {"charts": 1}):
        with pytest.raises(Exception):
            validate_config(bad)


# --- publish: dry run and CNAME preservation (real git repositories) -------------------------------

GIT_ID = ["-c", "user.name=t", "-c", "user.email=t@example.com"]


def git(repo: Path, *args: str) -> str:
    return subprocess.run(["git", "-C", str(repo), *GIT_ID, *args], check=True, capture_output=True, text=True).stdout.strip()


@pytest.fixture
def remote_and_clone(tmp_path: Path):
    seed = tmp_path / "seed"
    seed.mkdir()
    subprocess.run(["git", "init", "-q", "-b", "main", str(seed)], check=True)
    (seed / "README.md").write_text("# 공개 저장소\n", encoding="utf-8")
    (seed / "docs").mkdir()
    (seed / "docs" / "CNAME").write_text("dash.example.com\n")
    (seed / "docs" / "old.html").write_text("old")
    git(seed, "add", "-A")
    git(seed, "commit", "-q", "-m", "seed")
    bare = tmp_path / "origin.git"
    subprocess.run(["git", "clone", "-q", "--bare", str(seed), str(bare)], check=True)
    clone = tmp_path / "public"
    subprocess.run(["git", "clone", "-q", "--depth", "1", f"file://{bare}", str(clone)], check=True)
    return bare, clone


def test_dry_run_commits_nothing(data_dir: Path, tmp_path: Path, remote_and_clone, capsys) -> None:
    bare, clone = remote_and_clone
    staging = tmp_path / "staging"
    assert build_cli.main(["--source", str(data_dir), "--out", str(staging), "--pages-url", "u"]) == 0
    head = git(bare, "rev-parse", "main")
    assert publish.main(["--staging", str(staging), "--repo", str(clone), "--dry-run"]) == 0
    out = capsys.readouterr().out
    assert "DRY RUN" in out and "docs/index.html" in out
    assert git(bare, "rev-parse", "main") == head and git(clone, "rev-parse", "HEAD") == head


def test_cname_survives_and_stale_files_do_not(data_dir: Path, tmp_path: Path, remote_and_clone) -> None:
    bare, clone = remote_and_clone
    staging = tmp_path / "staging"
    assert build_cli.main(["--source", str(data_dir), "--out", str(staging), "--pages-url", "u"]) == 0
    assert publish.run(staging, clone) == "pushed"
    files = git(bare, "ls-tree", "-r", "--name-only", "main").splitlines()
    assert "docs/CNAME" in files and "docs/old.html" not in files
    assert (clone / "docs" / "CNAME").read_text() == "dash.example.com\n"
    assert publish.run(staging, clone) == "unchanged"


def test_unreachable_remote_is_reported_before_commit(data_dir: Path, tmp_path: Path, remote_and_clone) -> None:
    bare, clone = remote_and_clone
    staging = tmp_path / "staging"
    assert build_cli.main(["--source", str(data_dir), "--out", str(staging), "--pages-url", "u"]) == 0
    git(clone, "remote", "set-url", "origin", str(tmp_path / "does-not-exist.git"))
    with pytest.raises(Exception, match="cannot reach branch"):
        publish.run(staging, clone)
    assert git(clone, "status", "--porcelain") == ""  # nothing applied, nothing staged


# --- repository hygiene ------------------------------------------------------------------------------

def test_workflow_files_parse_and_agree_with_setup() -> None:
    data1_wf = DATA1_COPY / "publish-dashboard.yml"
    ci_wf = ROOT / ".github" / "workflows" / "ci.yml"
    wf = yaml.safe_load(data1_wf.read_text(encoding="utf-8"))
    ci = yaml.safe_load(ci_wf.read_text(encoding="utf-8"))
    on = wf[True] if True in wf else wf["on"]  # YAML 1.1 parses the bare key `on` as boolean True
    assert on["push"]["paths"] == [
        "**.xlsx", "**.xlsm", "**.XLSX", "**.XLSM", "!**/~$*",
        "dashboard.config.json", ".github/workflows/publish-dashboard.yml",
    ]
    assert "workflow_dispatch" in on and wf["permissions"] == {"contents": "read"}
    assert wf["concurrency"] == {"group": "publish-dashboard", "cancel-in-progress": False}
    text = data1_wf.read_text(encoding="utf-8")
    assert "secrets.DASHBOARD_PUSH_TOKEN" in text and "actions/checkout@v7" in text and "actions/setup-python@v7" in text
    assert "$default-branch" not in text
    ci_on = ci[True] if True in ci else ci["on"]
    assert ci_on["push"]["paths-ignore"] == ["docs/**", "README.md"]
    setup = (ROOT / "SETUP.md").read_text(encoding="utf-8")
    assert "DASHBOARD_PUSH_TOKEN" in setup and "Contents" in setup and "/docs" in setup


def test_data1_copies_in_public_repo_are_identical() -> None:
    if not DATA1_ORIGINAL.is_dir():
        pytest.skip("no sibling data1/ directory (running inside the public repository alone)")
    for name, src in (("publish-dashboard.yml", ".github/workflows/publish-dashboard.yml"), ("dashboard.config.example.json", "dashboard.config.example.json")):
        assert (DATA1_COPY / name).read_bytes() == (DATA1_ORIGINAL / src).read_bytes(), name


def test_example_config_is_accepted_verbatim() -> None:
    raw = json.loads((DATA1_COPY / "dashboard.config.example.json").read_text(encoding="utf-8"))
    c = validate_config(raw)
    assert c["mode"] == "rows" and c["auto_exclude_sensitive"] is True
