"""Wide daily time-series sheets: detection, extraction, cleaning, statistics, curves, site output."""

from __future__ import annotations

import datetime as dt
import itertools
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from dashboard_builder import build as build_cli
from dashboard_builder.config import validate_config
from dashboard_builder.excel import extract_workbook
from dashboard_builder.render_readme import fmt_ts_change, fmt_ts_value
from dashboard_builder.render_site import TEMPLATE_INDEX, TEMPLATE_MARKET
from dashboard_builder.sample import make_sample_workbook
from dashboard_builder.timeseries import (
    DETECT_ROWS,
    build_market,
    clean_series,
    detect_timeseries,
    display_name,
    infer_unit,
    parse_tenor,
    series_id,
)

from conftest import make_hard_workbook

HEADERS = ["한국_3Y", "한국_10Y", "한국_기준금리", "USDKRW", "US_IG_SP", "KTB_외국인순매수", "KDB_AAA_3Y", "KDB_AAA_5Y", "KDB_AAA_10Y"]
SPIKE_ROW = 50
INTERIOR_ZERO_ROW = 100
LEADING_ZEROS = 30


def business_days(start: dt.date, n: int):
    d = start
    while n:
        if d.weekday() < 5:
            yield d
            n -= 1
        d += dt.timedelta(days=1)


def make_market_workbook(path: Path, n_days: int = 400, descending: bool = False, plain_header: bool = False) -> list:
    """A market-data export: metadata rows, Category/Ticker/Notation rows, a DATES header, a field-label
    row, dated rows (one FX spike, leading zeros, one interior zero) and blank padding rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = "DailyRate"
    if not plain_header:
        ws.append(["시작", None, "종료", None, "Data 개수", 9999, "주기", "일"])
        ws.append([])
        ws.append(["Start Date", dt.datetime(2024, 1, 1)])
        ws.append(["End Date", dt.datetime(2025, 8, 1)])
        ws.append(["Category", "RATE", "RATE", "RATE", "FX", "RATE", "DEMAND", "CREDIT", "CREDIT", "CREDIT"])
        ws.append(["Ticker", "BONDAVG01", "BONDAVG01", "68", "환율_달러원", "미국_투자등급_스프레드", "한국_KTB_외국인순매수",
                   "한국_크레딧_특수채_AAA_15y", "한국_크레딧_특수채_AAA_15y", "한국_크레딧_특수채_AAA_15y"])  # copy-paste slip in the source
        ws.append(["Notation", "한국_3y", "한국_10y", "한국_기준금리", "USDKRW", "US_IG_SP", "KTB_외국인순매수", "KDB_AAA_3y", "KDB_AAA_5y", "KDB_AAA_10y"])
        ws.append(["DATES"] + HEADERS)
        ws.append(["일자", "3년이하(당일)", "10년이하(당일)", "현재가", "중간값", "OAS", "외국인 순매수거래대금", "3년이하(당일)", "5년이하(당일)", "10년이하(당일)"])
    else:
        ws.append(["일자"] + HEADERS)
    days = list(business_days(dt.date(2024, 1, 1), n_days))
    rows = []
    for k, day in enumerate(days):
        usd = 130000 if k == SPIKE_ROW else 1300 + k * 0.5
        rows.append([
            dt.datetime(day.year, day.month, day.day),
            round(3.0 + 0.001 * k, 4), round(3.3 + 0.001 * k, 4), 3.5 if k < 200 else 3.25, usd, 80 + (k % 7),
            (k % 5 - 2) * 1e9,
            0 if k < LEADING_ZEROS else round(3.4 + 0.001 * k, 4), round(3.6 + 0.001 * k, 4),
            0 if k == INTERIOR_ZERO_ROW else round(3.8 + 0.001 * k, 4),
        ])
    for row in (reversed(rows) if descending else rows):
        ws.append(row)
    for _ in range(20):
        ws.append([None] * 10)
    wb.save(path)
    return days


def head_rows(path: Path, sheet: str):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(itertools.islice(wb[sheet].iter_rows(values_only=True), DETECT_ROWS))
    finally:
        wb.close()


def market_of(path: Path, **overrides):
    cfg = validate_config(overrides)
    f = extract_workbook(path, path.name, cfg)
    ts = f.pop("timeseries")
    for t in ts:
        t["file"] = path.name
    return f, build_market(ts, cfg)


# --- detection ---------------------------------------------------------------------------------------

def test_detects_market_sheet_with_metadata_rows(tmp_path: Path) -> None:
    make_market_workbook(tmp_path / "m.xlsx")
    assert detect_timeseries(head_rows(tmp_path / "m.xlsx", "DailyRate")) == 7


def test_detects_plain_date_header_and_descending_order(tmp_path: Path) -> None:
    days = make_market_workbook(tmp_path / "p.xlsx", n_days=60, descending=True, plain_header=True)
    assert detect_timeseries(head_rows(tmp_path / "p.xlsx", "DailyRate")) == 0
    _, market = market_of(tmp_path / "p.xlsx")
    assert market["dates"] == [d.isoformat() for d in days]  # sorted ascending regardless of file order
    assert market["series"][0]["category"] == "기타" and market["series"][0]["name"] == "한국 3Y"


def test_transactional_sheets_are_not_time_series(tmp_path: Path) -> None:
    make_hard_workbook(tmp_path / "hard.xlsx")
    make_sample_workbook(tmp_path / "sample.xlsx")
    assert detect_timeseries(head_rows(tmp_path / "hard.xlsx", "매출현황")) is None
    assert detect_timeseries(head_rows(tmp_path / "sample.xlsx", "월별 매출")) is None  # text columns (지역, 제품) -> a table sorted by date
    assert detect_timeseries([["일자", "값"], [dt.datetime(2024, 1, 1), 1], [dt.datetime(2024, 1, 2), 2]]) is None  # too short


# --- naming, units, cleaning -------------------------------------------------------------------------

def test_display_name_prefers_korean_label_and_fixes_tenor_slip() -> None:
    assert display_name("UST10Y", "미국_10y", "UST10y") == "미국 10y"
    assert display_name("한국_3M", "BONDAVG14", "한국_3m") == "한국 3M"
    assert display_name("KDB_AAA_3Y", "한국_크레딧_특수채_AAA_15y", "KDB_AAA_3y") == "한국 크레딧 특수채 AAA 3Y"
    assert display_name("CAPITAL_AA_MINUS_1Y", "한국_크레딧_여전채_AA-_1y", "capital_AA_minus_1y") == "한국 크레딧 여전채 AA- 1y"
    assert parse_tenor("CAPITAL_AA_MINUS_1Y") == ("CAPITAL_AA_MINUS", 1.0, "1Y")
    assert parse_tenor("UST3M") == ("UST", 0.25, "3M") and parse_tenor("한국_1.5Y")[1:] == (1.5, "1.5Y")
    assert parse_tenor("USDKRW") is None
    assert series_id("A", "B", "C") == series_id("A", "B", "C") and series_id("A", "B", "C") != series_id("A", "B", "D")
    assert series_id("한국_3M", "BONDAVG14", "한국_3m").startswith("3m-")


def test_units() -> None:
    assert infer_unit("RATE", "UST10Y", "UST10y", "미국_10y", "MID_Close") == "pct"
    assert infer_unit("RATE", "US_IG_SP", "US_IG_SP", "미국_투자등급_스프레드", "OAS") == "bp"
    assert infer_unit("FX", "USDKRW", "USDKRW", "환율_달러원", "중간값") == "level"
    assert infer_unit("DEMAND", "KTB_외국인순매수", "", "", "") == "amount"
    assert infer_unit("MACRO", "VIX", "VIX", "미국_VIX", "현재가") == "pts"
    assert infer_unit("", "X", "", "", "") == "level"


def test_clean_series_rules() -> None:
    # leading zeros before a >= 0.5 rate are holes; near-zero policy rates keep their zeros
    vals, removed = clean_series([0, 0, 3.1, 3.2, 3.3, 3.2, 3.1, 3.0, 3.1], "pct")
    assert vals[:2] == [None, None] and removed == [0, 1]
    vals, removed = clean_series([0.05, 0.05, 0, 0, 0, 0, 0.5, 0.5], "pct")
    assert removed == [] and vals[2] == 0
    # interior zero between >= 0.5 neighbours is a hole
    vals, removed = clean_series([3.0, 3.1, 0, 3.1, 3.0, 3.1, 3.0], "pct")
    assert removed == [2] and vals[2] is None
    # a one-day spike among a consistent neighbourhood is dropped; a regime move is kept
    vals, removed = clean_series([1300, 1301, 1299, 1302, 1300, 130000, 1301, 1299, 1300, 1302, 1301], "level")
    assert removed == [5] and vals[5] is None
    vals, removed = clean_series([1.2, 1.2, 1.1, 1.3, 1.2, 0.07, 0.08, 0.1, 0.3, 0.5, 0.9], "pct")
    assert removed == []
    # zeros never count in price levels; flows are never touched
    vals, removed = clean_series([0.7, 0, 0.71, 0.7], "level")
    assert removed == [1]
    vals, removed = clean_series([0, 5e9, 0, -3e9], "amount")
    assert removed == [] and vals[0] == 0


# --- extraction + statistics -------------------------------------------------------------------------

def test_extract_and_build_market(tmp_path: Path) -> None:
    days = make_market_workbook(tmp_path / "m.xlsx")
    f, market = market_of(tmp_path / "m.xlsx")
    assert [s["state"] for s in f["sheets"]] == ["timeseries"]
    stub = f["sheets"][0]
    assert stub["rows_total"] == 400 and stub["columns_published"] == 9 and stub["rows"] == [] and "시계열" in stub["notes"][0]
    assert market["asof"] == days[-1].isoformat() and len(market["dates"]) == 400
    by_header = {s["header"]: s for s in market["series"]}
    assert list(by_header) == HEADERS
    assert [by_header[h]["name"] for h in HEADERS[:6]] == ["한국 3Y", "한국 10Y", "한국 기준금리", "환율 달러원", "미국 투자등급 스프레드", "한국 KTB 외국인순매수"]
    assert [by_header[h]["name"] for h in HEADERS[6:]] == ["한국 크레딧 특수채 AAA 3Y", "한국 크레딧 특수채 AAA 5Y", "한국 크레딧 특수채 AAA 10Y"]
    assert [by_header[h]["unit"] for h in HEADERS] == ["pct", "pct", "pct", "level", "bp", "amount", "pct", "pct", "pct"]
    assert by_header["한국_3Y"]["field"] == "3년이하(당일)" and by_header["US_IG_SP"]["decimals"] == 0 and by_header["한국_3Y"]["decimals"] == 3
    assert [c["key"] for c in market["categories"]] == ["RATE", "FX", "DEMAND", "CREDIT"]

    usd = by_header["USDKRW"]
    assert usd["cleaned"] == 1 and usd["cleaned_dates"] == [days[SPIKE_ROW].isoformat()]
    assert market["_values"][usd["id"]][SPIKE_ROW] is None and market["_values"][usd["id"]][SPIKE_ROW + 1] == 1300 + (SPIKE_ROW + 1) * 0.5
    kdb3 = by_header["KDB_AAA_3Y"]
    assert kdb3["cleaned"] == LEADING_ZEROS and kdb3["first"] == LEADING_ZEROS
    kdb10 = by_header["KDB_AAA_10Y"]
    assert kdb10["cleaned"] == 1 and market["_values"][kdb10["id"]][INTERIOR_ZERO_ROW] is None
    assert market["cleaning"] == {"enabled": True, "points": LEADING_ZEROS + 2}

    kr3 = by_header["한국_3Y"]
    last = len(days) - 1
    assert kr3["latest"] == {"i": last, "date": days[-1].isoformat(), "value": round(3.0 + 0.001 * last, 4)}
    assert kr3["bases"]["d1"]["i"] == last - 1
    assert kr3["bases"]["w1"]["date"] <= (days[-1] - dt.timedelta(days=7)).isoformat()
    assert kr3["bases"]["ytd"]["date"] == max(d for d in days if d.year == days[-1].year - 1).isoformat()
    assert kr3["bases"]["y1"]["date"] <= dt.date(days[-1].year - 1, days[-1].month, days[-1].day).isoformat()
    assert kr3["range52"]["hi"]["i"] == last and kr3["extremes"]["lo"]["i"] == 0
    assert 5 <= len(kr3["spark"]) <= 53 and kr3["spark"][-1][0] == last
    assert "sums" not in kr3
    flow = by_header["KTB_외국인순매수"]
    vals = market["_values"][flow["id"]]
    assert flow["sums"]["d1"] == vals[-1] and flow["sums"]["d5"] == sum(vals[-5:]) and "bases" not in flow
    assert flow["sums"]["ytd"] == sum(v for d, v in zip(days, vals) if d.year == days[-1].year)

    assert [c["label"] for c in market["curves"]] == ["한국 크레딧 특수채 AAA"]  # 한국_3Y/10Y are only two tenors
    curve = market["curves"][0]
    assert [t["label"] for t in curve["tenors"]] == ["3Y", "5Y", "10Y"] and curve["unit"] == "pct"
    assert [s["key"] for s in curve["snapshots"]] == ["latest", "w1", "m1", "y1"]
    assert curve["snapshots"][0]["date"] == days[-1].isoformat() and curve["snapshots"][0]["values"][1] == round(3.6 + 0.001 * last, 4)

    assert len(market["highlights"]) == 8 and market["highlights"][0] == kr3["id"] and market["highlights"][1] == usd["id"]  # round-robin over categories
    assert market["featured"] and all(g["series"] for g in market["featured"])
    assert market["notes"] == []


def test_config_highlights_featured_and_column_filters(tmp_path: Path) -> None:
    make_market_workbook(tmp_path / "m.xlsx")
    _, market = market_of(
        tmp_path / "m.xlsx",
        exclude_columns=["KDB_*"],
        timeseries={
            "highlights": ["ust10y", "usdkrw", "한국 3y"],
            "featured": [{"title": "환율", "series": ["USDKRW", "US_IG_SP"], "mode": "rebase"}, ["nothing"]],
        },
    )
    headers = [s["header"] for s in market["series"]]
    assert not any(h.startswith("KDB") for h in headers) and market["sources"][0]["notes"][0].startswith("설정으로 제외한 지표 3개")
    names = {s["id"]: s["header"] for s in market["series"]}
    assert [names[i] for i in market["highlights"]] == ["USDKRW", "한국_3Y"]
    assert market["featured"] == [{"title": "환율", "series": [s["id"] for s in market["series"] if s["header"] in ("USDKRW", "US_IG_SP")], "mode": "rebase"}]
    assert any("ust10y" in n for n in market["notes"]) and any("nothing" in n for n in market["notes"])


def test_timeseries_config_validation() -> None:
    cfg = validate_config({"timeseries": {"enabled": False, "highlights": ["a"], "featured": [["a", "b"], {"title": "t", "series": ["a"], "mode": "diff"}], "spark_points": 10}})
    assert cfg["timeseries"]["enabled"] is False and cfg["timeseries"]["max_series"] == 1000
    for bad in (
        {"timeseries": {"typo": 1}},
        {"timeseries": {"featured": [[]]}},
        {"timeseries": {"featured": [{"series": ["a"], "mode": "log"}]}},
        {"timeseries": {"highlights": "a"}},
        {"timeseries": {"spark_points": 1}},
        {"timeseries": []},
    ):
        with pytest.raises(Exception):
            validate_config(bad)


def test_readme_formatting_helpers() -> None:
    assert fmt_ts_value(4.0399, "pct", 2) == "4.04%" and fmt_ts_value(78.0, "bp", 0) == "78bp"
    assert fmt_ts_value(1377.15, "level", 2) == "1,377.15" and fmt_ts_value(2053890035.0, "amount", 0) == "2,053,890,035"
    assert fmt_ts_change(4.04, 4.10, "pct") == "-6.0bp" and fmt_ts_change(78, 80, "bp") == "-2.0bp"
    assert fmt_ts_change(1377.15, 1364.5, "level") == "+0.93%" and fmt_ts_change(44.4, 45.63, "pts") == "-1.23"
    assert fmt_ts_change(None, 1, "pct") == "–" and fmt_ts_change(1, 0, "level") == "–"


# --- end to end --------------------------------------------------------------------------------------

def _tree(path: Path):
    return {p.relative_to(path).as_posix(): p.read_bytes() for p in sorted(path.rglob("*")) if p.is_file()}


def test_build_writes_market_site_and_is_idempotent(tmp_path: Path, capsys) -> None:
    src = tmp_path / "data1"
    src.mkdir()
    make_market_workbook(src / "rates.xlsx")
    make_hard_workbook(src / "tables.xlsx")
    (src / "dashboard.config.json").write_text(json.dumps({"title": "마켓", "timeseries": {"highlights": ["USDKRW", "한국_10Y"]}}, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "staging"
    assert build_cli.main(["--source", str(src), "--out", str(out), "--pages-url", "https://tus4.github.io/test/"]) == 0
    printed = capsys.readouterr().out
    assert "Time-series dashboard: 9 series, 400 dates" in printed and "rates.xlsx / DailyRate: timeseries" in printed

    docs = out / "docs"
    assert (docs / "index.html").read_bytes() == TEMPLATE_MARKET.read_bytes()
    assert (docs / "tables.html").read_bytes() == TEMPLATE_INDEX.read_bytes()
    market = json.loads((docs / "data" / "market.json").read_text(encoding="utf-8"))
    assert market["title"] == "마켓" and len(market["series"]) == 9 and market["tables_page"] == "tables.html"
    assert "_values" not in market and len(market["dates"]) == 400
    series_dir = docs / "data" / "series"
    assert sorted(p.name for p in series_dir.iterdir()) == sorted(s["id"] + ".json" for s in market["series"])
    usd = next(s for s in market["series"] if s["header"] == "USDKRW")
    payload = json.loads((series_dir / (usd["id"] + ".json")).read_text(encoding="utf-8"))
    assert payload["start"] == usd["first"] and len(payload["values"]) == usd["last"] - usd["first"] + 1 and payload["values"][SPIKE_ROW] is None

    data = json.loads((docs / "data" / "dashboard.json").read_text(encoding="utf-8"))
    states = {(f["path"], s["name"]): s["state"] for f in data["files"] for s in f["sheets"]}
    assert states[("rates.xlsx", "DailyRate")] == "timeseries" and states[("tables.xlsx", "매출현황")] == "ok"
    assert data["timeseries"]["series_count"] == 9 and [h["header"] if "header" in h else h["name"] for h in data["timeseries"]["highlights"]] == ["환율 달러원", "한국 10Y"]
    assert not any("timeseries" in f for f in data["files"])  # the raw series never enter dashboard.json
    block = (out / "readme_block.md").read_text(encoding="utf-8")
    assert "### 📈 시계열 지표" in block and "| 환율 달러원 |" in block and "시계열 데이터입니다" in block
    summary = json.loads((out / "build_summary.json").read_text(encoding="utf-8"))
    assert summary["timeseries"]["series"] == 9 and summary["timeseries"]["curves"] == 1

    first = _tree(out)
    assert build_cli.main(["--source", str(src), "--out", str(out), "--pages-url", "https://tus4.github.io/test/"]) == 0
    assert _tree(out) == first  # byte-identical on unchanged input


def test_timeseries_can_be_disabled(tmp_path: Path) -> None:
    src = tmp_path / "data1"
    src.mkdir()
    make_market_workbook(src / "rates.xlsx", n_days=30)
    (src / "dashboard.config.json").write_text('{"timeseries": {"enabled": false}}', encoding="utf-8")
    out = tmp_path / "staging"
    assert build_cli.main(["--source", str(src), "--out", str(out)]) == 0
    docs = out / "docs"
    assert (docs / "index.html").read_bytes() == TEMPLATE_INDEX.read_bytes() and not (docs / "tables.html").exists()
    assert not (docs / "data" / "market.json").exists()
    data = json.loads((docs / "data" / "dashboard.json").read_text(encoding="utf-8"))
    assert data["timeseries"] is None and data["files"][0]["sheets"][0]["state"] == "ok"
